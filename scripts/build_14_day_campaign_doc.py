from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/home/ubuntu/promptforge')
ASSET_ROOT = Path('/home/ubuntu/webdev-static-assets/promptforge-16x9')
FRAMEWORK = ROOT / 'docs/14-day-campaign-framework.md'
CONTENT = ROOT / 'docs/14-day-campaign-content.md'
OUT = Path('/home/ubuntu/PromptForge_14-Day_Social_Media_Campaign_Calendar.docx')

DAY_IMAGES = {
    1: ['A-homepage-hero-16x9.png'],
    2: ['B-sign-up-page-16x9.png'],
    3: ['C-builder-empty-state-16x9.png'],
    4: ['D-social-media-selected-16x9.png'],
    5: ['E-live-prompt-mid-build-16x9.png'],
    6: ['F-complete-forged-prompt-16x9.png'],
    7: ['G-copy-save-controls-16x9.png'],
    8: ['H-prompt-library-grid-16x9.png'],
    9: ['I-unlock-pricing-16x9.png'],
    10: ['J-contact-us-page-16x9.png'],
    11: ['C-builder-empty-state-16x9.png', 'F-complete-forged-prompt-16x9.png'],
    12: ['G-copy-save-controls-16x9.png', 'H-prompt-library-grid-16x9.png'],
    13: ['D-social-media-selected-16x9.png', 'E-live-prompt-mid-build-16x9.png'],
    14: ['A-homepage-hero-16x9.png', 'F-complete-forged-prompt-16x9.png', 'I-unlock-pricing-16x9.png'],
}

ORANGE = RGBColor(237, 112, 36)
DARK = RGBColor(33, 33, 33)
STEEL = RGBColor(92, 102, 112)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'E56F25')
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_markdown_block(doc, lines):
    in_quote = False
    for raw in lines:
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith('> '):
            p = doc.add_paragraph(style='Intense Quote')
            p.paragraph_format.space_after = Pt(5)
            p.add_run(line[2:])
            continue
        if line.startswith('**') and line.endswith('**'):
            p = doc.add_paragraph()
            r = p.add_run(line.strip('*'))
            r.bold = True
            r.font.color.rgb = ORANGE
            continue
        if line.startswith('**') and ':' in line:
            p = doc.add_paragraph()
            label, rest = line.split(':', 1)
            r = p.add_run(label + ':')
            r.bold = True
            r.font.color.rgb = ORANGE
            p.add_run(rest)
            continue
        if line.startswith('---'):
            add_hr(doc)
            continue
        if line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(line[2:])
            continue
        if line.startswith('|'):
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        # Simple emphasis conversion.
        parts = re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.bold = True
            else:
                p.add_run(part)


def add_screenshot_block(doc, day):
    images = DAY_IMAGES[day]
    if len(images) == 1:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(ASSET_ROOT / images[0]), width=Inches(6.25))
        cap = doc.add_paragraph(f'Day {day} visual: PromptForge website screenshot.', style='Caption')
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        table = doc.add_table(rows=1, cols=len(images))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for idx, image in enumerate(images):
            cell = table.cell(0, idx)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ASSET_ROOT / image), width=Inches(6.05 / len(images)))
        cap = doc.add_paragraph(f'Day {day} visual sequence: matching PromptForge product screenshots.', style='Caption')
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def extract_day_sections(text):
    matches = list(re.finditer(r'^## Day (\d+) —.*$', text, re.MULTILINE))
    sections = {}
    for i, match in enumerate(matches):
        day = int(match.group(1))
        start = match.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else text.find('\n# Paid advertising scripts')
        sections[day] = text[start:end].strip().splitlines()
    return sections


def extract_after_heading(text, heading):
    marker = heading + '\n'
    start = text.find(marker)
    if start < 0:
        return []
    start += len(marker)
    next_heading = text.find('\n## ', start)
    return text[start: next_heading if next_heading >= 0 else len(text)].strip().splitlines()


def configure_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal'].font.size = Pt(9.5)
    styles['Title'].font.name = 'Aptos Display'
    styles['Title'].font.size = Pt(28)
    styles['Title'].font.color.rgb = DARK
    styles['Heading 1'].font.name = 'Aptos Display'
    styles['Heading 1'].font.size = Pt(18)
    styles['Heading 1'].font.color.rgb = ORANGE
    styles['Heading 2'].font.name = 'Aptos Display'
    styles['Heading 2'].font.size = Pt(13)
    styles['Heading 2'].font.color.rgb = DARK


def main():
    framework_text = FRAMEWORK.read_text(encoding='utf-8')
    content_text = CONTENT.read_text(encoding='utf-8')
    day_sections = extract_day_sections(content_text)
    doc = Document()
    configure_doc(doc)

    title = doc.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run('PromptForge 14-Day Social Media Campaign Calendar')
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run('A story-led growth campaign for LinkedIn, Facebook, Instagram, TikTok, and X')
    r.italic = True
    r.font.color.rgb = STEEL
    doc.add_paragraph('Prepared for PromptForge | Workshop Noir campaign system | Author: Manus AI', style='Subtitle')
    add_hr(doc)

    doc.add_heading('Campaign at a glance', level=1)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for cell, text in zip(table.rows[0].cells, ['Campaign element', 'Plan']):
        set_cell_text(cell, text, bold=True, color=RGBColor(255, 255, 255))
        set_cell_shading(cell, 'E56F25')
    overview = [
        ('Primary objective', 'Increase qualified visits, Builder starts, sign-ups, Library exploration, and WhatsApp unlock conversations.'),
        ('Core promise', 'Move from a rough idea to a clearer, structured AI prompt.'),
        ('Campaign idea', 'From Blank Page to Forged Prompt: A 14-Day PromptForge Workshop.'),
        ('Primary CTA', 'Start free with the PromptForge Builder.'),
        ('Paid CTA', 'Visit Pricing and message the main PromptForge WhatsApp Business chat.'),
        ('Offer shown', 'Lifetime access: ₦10,000 / $10.'),
    ]
    for key, value in overview:
        cells = table.add_row().cells
        set_cell_text(cells[0], key, bold=True, color=DARK)
        set_cell_text(cells[1], value)

    doc.add_heading('Campaign framework', level=1)
    framework_lines = framework_text.splitlines()
    add_markdown_block(doc, framework_lines[2:])

    doc.add_page_break()
    doc.add_heading('Fourteen-day publishing calendar', level=1)
    intro = doc.add_paragraph('Each daily section below is ready to adapt across the five requested platforms. The screenshots are placed beside the corresponding story so the document can be handed directly to a social media manager or creator.')
    intro.paragraph_format.space_after = Pt(8)

    for day in range(1, 15):
        doc.add_heading(f'Day {day}', level=1)
        add_screenshot_block(doc, day)
        add_markdown_block(doc, day_sections[day])
        if day in (4, 8, 12):
            doc.add_page_break()

    # Append the advertising and operational sections from the content file.
    paid_start = content_text.find('# Paid advertising scripts')
    if paid_start >= 0:
        doc.add_page_break()
        doc.add_heading('Paid advertising scripts and optimization guidance', level=1)
        add_markdown_block(doc, content_text[paid_start:].splitlines()[1:])

    doc.add_page_break()
    doc.add_heading('Screenshot asset index', level=1)
    doc.add_paragraph('The daily visuals in this document use the following supplied PromptForge website screenshots. The original image files remain available as separate assets for editing, scheduling, or creator handoff.')
    asset_table = doc.add_table(rows=1, cols=3)
    asset_table.style = 'Table Grid'
    asset_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(asset_table.rows[0].cells, ['Label', 'Screenshot', 'Primary use']):
        set_cell_text(cell, text, bold=True, color=RGBColor(255, 255, 255))
        set_cell_shading(cell, 'E56F25')
    uses = {
        'A': 'Homepage Hero', 'B': 'Sign-up Page', 'C': 'Builder Empty State', 'D': 'Social Media Selected',
        'E': 'Live Prompt Mid-Build', 'F': 'Complete Forged Prompt', 'G': 'Copy/Save Controls',
        'H': 'Prompt Library Grid', 'I': 'Unlock/Pricing', 'J': 'Contact Page'
    }
    for label, name in uses.items():
        row = asset_table.add_row().cells
        set_cell_text(row[0], label, bold=True, color=ORANGE)
        set_cell_text(row[1], name)
        set_cell_text(row[2], 'Use as the product proof image for the matching daily theme.')

    doc.add_paragraph('Production note: use the 16:9 versions for LinkedIn, Facebook, and X; use the 9:16 versions for TikTok and Instagram Reels. Keep the product UI recognizable and place only short overlay headlines on the screenshots.', style='Intense Quote')
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
